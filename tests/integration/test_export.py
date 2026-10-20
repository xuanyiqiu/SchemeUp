import json
import re
from dataclasses import replace
from datetime import datetime
from pathlib import Path

import pytest
from PySide6.QtGui import QImage
from PySide6.QtWidgets import QApplication

from schemeup.application.export import ExportFormat, ExportRequest, ExportScope
from schemeup.application.library import fit_document
from schemeup.application.ports import MemoryAssetStore
from schemeup.domain import commands
from schemeup.domain.commands import ShapeKind
from schemeup.domain.document import CanvasSettings, Document
from schemeup.domain.elements import Arrow, Frame, Rectangle
from schemeup.domain.errors import ExportError, UnsupportedFormatError
from schemeup.domain.export_layout import (
    PageLayout,
    export_regions,
    fit_scale,
    is_infinite,
    page_rects,
    tiles,
)
from schemeup.domain.geometry import Point, Rect, Size
from schemeup.domain.style import Style
from schemeup.domain.templates import architecture
from schemeup.infrastructure.export import FileExporter
from schemeup.infrastructure.formats import FileImporter
from schemeup.rendering.fonts import register_fonts
from schemeup.rendering.text import QtTextMeasurer


def frame(rect: Rect, name: str) -> Frame:
    element = commands.make_shape(ShapeKind.FRAME, rect, Style())
    assert isinstance(element, Frame)
    return replace(element, name=name)


@pytest.fixture
def document(qapp: QApplication) -> Document:
    register_fonts()
    measurer = QtTextMeasurer()
    base = fit_document(architecture(datetime(2026, 10, 20)), measurer)
    frame_a = frame(Rect(-40, 80, 600, 180), "Path")
    frame_b = frame(Rect(600, 80, 300, 180), "Data")
    return base.with_elements({frame_a.id: frame_a, frame_b.id: frame_b})


@pytest.fixture
def exporter() -> FileExporter:
    return FileExporter(MemoryAssetStore())


def test_layout_helpers() -> None:
    assert len(page_rects("A4", False, 2000)) == 2  # A4 is 1122 units tall
    assert len(page_rects("A4", False, 2400)) == 3
    assert page_rects("A4", False, 0)[0].left < 0  # the origin sits inside the margin
    assert page_rects(None, False, 2000) == []
    assert len(tiles(Rect(0, 0, 250, 150), Size(100, 100))) == 6
    assert fit_scale(Rect(0, 0, 200, 100), Size(100, 100)) == 0.5
    assert is_infinite(Document())
    assert not is_infinite(Document(settings=CanvasSettings(page_size="A4")))


def test_regions_follow_the_layout(document: Document) -> None:
    assert len(export_regions(document, PageLayout.FRAMES, None, "A4", False)) == 2
    assert export_regions(document, PageLayout.FIT, Rect(0, 0, 10, 10), "A4", False) == [
        Rect(0, 0, 10, 10)
    ]


def test_png_and_svg(document: Document, exporter: FileExporter, tmp_path: Path) -> None:
    png, svg = tmp_path / "a.png", tmp_path / "a.svg"
    exporter.export(document, ExportRequest(ExportFormat.PNG, scale=1), frozenset(), png)
    exporter.export(document, ExportRequest(ExportFormat.SVG), frozenset(), svg)
    image = QImage(str(png))
    assert image.width() > 500 and image.height() > 100
    text = svg.read_text(encoding="utf-8")
    assert "<svg" in text and "<text" not in text  # text became outlines


def test_selection_only_export_is_smaller(
    document: Document, exporter: FileExporter, tmp_path: Path
) -> None:
    shape = next(e for e in document.ordered() if isinstance(e, Rectangle))
    full, part = tmp_path / "full.png", tmp_path / "part.png"
    exporter.export(document, ExportRequest(ExportFormat.PNG, scale=1), frozenset(), full)
    exporter.export(
        document,
        ExportRequest(ExportFormat.PNG, scope=ExportScope.SELECTION, scale=1),
        frozenset({shape.id}),
        part,
    )
    assert QImage(str(part)).width() < QImage(str(full)).width()


def test_pdf_has_one_page_per_frame(
    document: Document, exporter: FileExporter, tmp_path: Path
) -> None:
    path = tmp_path / "a.pdf"
    exporter.export(
        document, ExportRequest(ExportFormat.PDF, layout=PageLayout.FRAMES), frozenset(), path
    )
    assert len(re.findall(rb"/Type\s*/Page[^s]", path.read_bytes())) == 2


def test_docx_text_notes_stay_editable(
    exporter: FileExporter, tmp_path: Path, qapp: QApplication
) -> None:
    from docx import Document as WordDocument

    text = commands.make_text(Point(0, 0), "Standup\n- ship it", Style())
    path = tmp_path / "notes.docx"
    exporter.export(
        Document.of([replace(text, width=200, height=50)]),
        ExportRequest(ExportFormat.DOCX, title="Standup"),
        frozenset(),
        path,
    )
    word = WordDocument(str(path))
    assert len(word.inline_shapes) == 0
    assert "- ship it" in [p.text for p in word.paragraphs]


def test_docx_drawings_become_images(
    document: Document, exporter: FileExporter, tmp_path: Path
) -> None:
    from docx import Document as WordDocument

    path = tmp_path / "sketch.docx"
    exporter.export(
        document, ExportRequest(ExportFormat.DOCX, layout=PageLayout.FRAMES), frozenset(), path
    )
    assert len(WordDocument(str(path)).inline_shapes) == 2


def test_empty_note_cannot_be_exported(exporter: FileExporter, tmp_path: Path) -> None:
    with pytest.raises(ExportError):
        exporter.export(
            Document(), ExportRequest(ExportFormat.PNG), frozenset(), tmp_path / "x.png"
        )


def test_schemeup_file_round_trip(
    document: Document, exporter: FileExporter, tmp_path: Path
) -> None:
    path = tmp_path / "note.schemeup"
    exporter.export(
        document,
        ExportRequest(ExportFormat.SCHEMEUP, title="Cache", tags=("arch",)),
        frozenset(),
        path,
    )
    imported = FileImporter().read(path)
    assert imported.title == "Cache" and imported.tags == ("arch",)
    assert imported.document.ordered() == document.ordered()


def test_excalidraw_import(tmp_path: Path) -> None:
    data = {
        "type": "excalidraw",
        "elements": [
            {
                "id": "r",
                "type": "rectangle",
                "x": 0,
                "y": 0,
                "width": 100,
                "height": 50,
                "roundness": None,
            },
            {
                "id": "t",
                "type": "text",
                "x": 10,
                "y": 10,
                "text": "API",
                "containerId": "r",
                "fontFamily": 3,
            },
            {
                "id": "a",
                "type": "arrow",
                "x": 100,
                "y": 25,
                "points": [[0, 0], [80, 0]],
                "endArrowhead": "triangle",
                "endBinding": {"elementId": "r", "gap": 4},
            },
            {"id": "d", "type": "ellipse", "x": 0, "y": 0, "isDeleted": True},
        ],
        "appState": {"viewBackgroundColor": "#fff9db"},
    }
    path = tmp_path / "sketch.excalidraw"
    path.write_text(json.dumps(data), encoding="utf-8")
    note = FileImporter().read(path)
    rectangle, arrow = note.document.ordered()
    assert isinstance(rectangle, Rectangle) and rectangle.label == "API"
    assert rectangle.style.roundness.value == "sharp"
    assert isinstance(arrow, Arrow) and arrow.end_binding is not None
    assert note.document.settings.background == "#fff9db"


def test_unknown_files_are_refused(tmp_path: Path) -> None:
    path = tmp_path / "x.json"
    path.write_text('{"hello": 1}', encoding="utf-8")
    with pytest.raises(UnsupportedFormatError):
        FileImporter().read(path)
